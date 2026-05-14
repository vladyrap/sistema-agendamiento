"""Generador de archivos Excel (.xlsx) para exports del sistema.

Uso típico:
    rows = [{"col1": "a", "col2": 1}, ...]
    headers = [{"key": "col1", "label": "Columna 1"}, {"key": "col2", "label": "Columna 2"}]
    return excel_response(rows, headers, "mis_datos")
"""
from __future__ import annotations
import io
from datetime import datetime, date, time
from typing import Any, Iterable, Optional
from fastapi.responses import StreamingResponse


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _coerce(v: Any) -> Any:
    """Convierte valores no-Excel-friendly a tipos compatibles."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date, time)):
        return v
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, (list, tuple, dict, set)):
        return str(v)
    return v


def excel_response(
    rows: Iterable[dict],
    headers: list[dict],   # [{"key": str, "label": str, "width": int (optional)}, ...]
    filename: str,
    sheet_name: str = "Datos",
    title: Optional[str] = None,
    subtitle: Optional[str] = None,
) -> StreamingResponse:
    """Crea un .xlsx con `openpyxl` y lo devuelve como respuesta HTTP."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # excel limita a 31 chars

    # Estilos
    title_font = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_cursor = 1
    if title:
        ws.cell(row=row_cursor, column=1, value=title).font = title_font
        ws.merge_cells(start_row=row_cursor, end_row=row_cursor,
                       start_column=1, end_column=max(1, len(headers)))
        row_cursor += 1
    if subtitle:
        ws.cell(row=row_cursor, column=1, value=subtitle).font = subtitle_font
        ws.merge_cells(start_row=row_cursor, end_row=row_cursor,
                       start_column=1, end_column=max(1, len(headers)))
        row_cursor += 1
    if title or subtitle:
        row_cursor += 1  # línea en blanco

    # Encabezados
    header_row = row_cursor
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col_idx, value=h["label"])
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = cell_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = h.get("width", 18)
    ws.row_dimensions[header_row].height = 22
    row_cursor += 1

    # Datos
    for r in rows:
        for col_idx, h in enumerate(headers, start=1):
            value = _coerce(r.get(h["key"]))
            c = ws.cell(row=row_cursor, column=col_idx, value=value)
            c.border = cell_border
            # Formato fecha/hora
            if isinstance(value, datetime):
                c.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                c.number_format = "yyyy-mm-dd"
            elif isinstance(value, time):
                c.number_format = "hh:mm"
        row_cursor += 1

    # Congelar primera fila de encabezado
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # Auto-filter sobre el rango
    if row_cursor > header_row + 1:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A{header_row}:{last_col}{row_cursor - 1}"

    # Stream en memoria
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = filename.replace('"', '').replace('\n', '_')
    if not safe_name.endswith(".xlsx"):
        safe_name += ".xlsx"

    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{safe_name}"',
            "Cache-Control": "no-store",
        },
    )
